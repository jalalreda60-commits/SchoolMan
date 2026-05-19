import os
from datetime import datetime
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                                QFrame, QComboBox, QSpinBox, QMessageBox, QTableWidget,
                                QTableWidgetItem, QHeaderView, QAbstractItemView)
from PySide6.QtCore import Qt

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

MONTHS = ['Janvier','Février','Mars','Avril','Mai','Juin',
          'Juillet','Août','Septembre','Octobre','Novembre','Décembre']


class ReportsWidget(QWidget):
    def __init__(self, session, current_user):
        super().__init__()
        self.session = session
        self.current_user = current_user
        self.setup_ui()

    def setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 15, 20, 15)
        layout.setSpacing(15)

        title = QLabel("📊 Rapports & Exports")
        title.setStyleSheet("color: #e0e0e0; font-size: 22px; font-weight: bold;")
        layout.addWidget(title)

        # Report cards grid
        cards_layout = QHBoxLayout()
        cards_layout.setSpacing(12)

        reports = [
            ("📋 Liste des Élèves", "Exporter la liste complète des élèves", "#3498db", self.report_students),
            ("💰 Rapport Financier", "Revenus, dépenses et bilan mensuel", "#27ae60", self.report_financial),
            ("⚠ Élèves Impayés", "Liste des élèves avec paiements en retard", "#e67e22", self.report_unpaid),
            ("👥 Rapport Employés", "Liste des employés et salaires", "#9b59b6", self.report_employees),
        ]

        for title_txt, desc, color, func in reports:
            card = QFrame()
            card.setObjectName("card")
            card.setStyleSheet(f"""
                QFrame#card {{
                    background: #1a1a2e;
                    border: 1px solid #2a2a4e;
                    border-top: 3px solid {color};
                    border-radius: 12px;
                }}
                QFrame#card:hover {{ background: #1e1e38; }}
            """)
            card_layout = QVBoxLayout(card)
            card_layout.setContentsMargins(20, 20, 20, 20)
            card_layout.setSpacing(8)

            ttl = QLabel(title_txt)
            ttl.setStyleSheet(f"color: {color}; font-size: 14px; font-weight: bold; background: transparent;")
            card_layout.addWidget(ttl)

            desc_lbl = QLabel(desc)
            desc_lbl.setStyleSheet("color: #808080; font-size: 11px; background: transparent;")
            desc_lbl.setWordWrap(True)
            card_layout.addWidget(desc_lbl)

            btn = QPushButton("📥 Générer PDF")
            btn.setStyleSheet(f"""
                QPushButton {{ background: {color}; color: white; border: none; border-radius: 6px;
                              padding: 8px; font-weight: bold; }}
                QPushButton:hover {{ opacity: 0.8; }}
            """)
            btn.clicked.connect(func)
            card_layout.addWidget(btn)
            cards_layout.addWidget(card)

        layout.addLayout(cards_layout)

        # Filters
        filter_frame = QFrame()
        filter_frame.setObjectName("card")
        filter_layout = QHBoxLayout(filter_frame)
        filter_layout.setContentsMargins(15, 10, 15, 10)

        filter_layout.addWidget(QLabel("Mois:"))
        self.month_combo = QComboBox()
        self.month_combo.setFixedHeight(36)
        self.month_combo.setFixedWidth(120)
        for m in MONTHS:
            self.month_combo.addItem(m)
        self.month_combo.setCurrentIndex(datetime.now().month - 1)
        filter_layout.addWidget(self.month_combo)

        filter_layout.addWidget(QLabel("Année:"))
        self.year_spin = QSpinBox()
        self.year_spin.setRange(2020, 2050)
        self.year_spin.setValue(datetime.now().year)
        self.year_spin.setFixedHeight(36)
        self.year_spin.setFixedWidth(80)
        filter_layout.addWidget(self.year_spin)

        filter_layout.addStretch()

        excel_btn = QPushButton("📊 Exporter Excel")
        excel_btn.setObjectName("successBtn")
        excel_btn.clicked.connect(self.export_excel)
        filter_layout.addWidget(excel_btn)

        layout.addWidget(filter_frame)

        # Preview table
        preview_lbl = QLabel("📋 Aperçu des données")
        preview_lbl.setStyleSheet("color: #FF6B00; font-size: 14px; font-weight: bold;")
        layout.addWidget(preview_lbl)

        self.preview_table = QTableWidget()
        self.preview_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.preview_table.verticalHeader().setVisible(False)
        self.preview_table.setShowGrid(False)
        self.preview_table.setAlternatingRowColors(True)
        layout.addWidget(self.preview_table)

        self.load_preview()

    def load_preview(self):
        from models.database import Payment, Student
        month = self.month_combo.currentIndex() + 1
        year = self.year_spin.value()
        payments = self.session.query(Payment).filter_by(month=month, year=year).all()

        self.preview_table.setColumnCount(5)
        self.preview_table.setHorizontalHeaderLabels(["Élève", "Type", "Montant", "Reçu", "Date"])
        self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.preview_table.setRowCount(0)

        type_labels = {'monthly': 'Mensualité', 'insurance': 'Assurance', 'transport': 'Transport'}
        for p in payments:
            student = p.student
            row = self.preview_table.rowCount()
            self.preview_table.insertRow(row)
            items = [
                f"{student.first_name} {student.last_name}" if student else "—",
                type_labels.get(p.payment_type, p.payment_type),
                f"{p.amount:,.2f} MAD",
                p.receipt_number or "—",
                p.payment_date.strftime("%d/%m/%Y") if p.payment_date else "—"
            ]
            for col, text in enumerate(items):
                item = QTableWidgetItem(text)
                item.setTextAlignment(Qt.AlignCenter)
                self.preview_table.setItem(row, col, item)
            self.preview_table.setRowHeight(row, 36)

    def report_students(self):
        self._generate_students_pdf()

    def report_financial(self):
        self._generate_financial_pdf()

    def report_unpaid(self):
        self._generate_unpaid_pdf()

    def report_employees(self):
        self._generate_employees_pdf()

    def _generate_students_pdf(self):
        from models.database import Student
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.lib.enums import TA_CENTER

        reports_dir = os.path.join(BASE_DIR, 'reports')
        os.makedirs(reports_dir, exist_ok=True)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filepath = os.path.join(reports_dir, f'students_{timestamp}.pdf')

        doc = SimpleDocTemplate(filepath, pagesize=landscape(A4),
                               rightMargin=15*mm, leftMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle('title', fontSize=16, fontName='Helvetica-Bold',
                                      textColor=colors.HexColor('#FF6B00'), alignment=TA_CENTER)
        story.append(Paragraph("LISTE DES ÉLÈVES", title_style))
        story.append(Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        story.append(Spacer(1, 10*mm))

        students = self.session.query(Student).filter_by(is_active=True).all()
        data = [["Code", "Nom", "Prénom", "Classe", "Parent", "Téléphone", "Transport", "Assurance"]]
        for s in students:
            data.append([
                s.student_code or "", s.last_name, s.first_name,
                s.class_.name if s.class_ else "—",
                s.parent_name or "—", s.parent_phone or "—",
                "Oui" if s.has_transport else "Non",
                "Oui" if s.insurance_paid else "Non"
            ])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#FF6B00')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 9),
            ('FONTSIZE', (0,1), (-1,-1), 8),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#FFF5EE')]),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e0e0e0')),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(table)
        doc.build(story)
        self._open_file(filepath)
        QMessageBox.information(self, "Succès", f"✅ Rapport généré:\n{filepath}")

    def _generate_financial_pdf(self):
        from models.database import Payment, Expense, Setting
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.lib.enums import TA_CENTER
        from sqlalchemy import func

        reports_dir = os.path.join(BASE_DIR, 'reports')
        os.makedirs(reports_dir, exist_ok=True)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filepath = os.path.join(reports_dir, f'financial_{timestamp}.pdf')

        year = self.year_spin.value()
        doc = SimpleDocTemplate(filepath, pagesize=A4,
                               rightMargin=20*mm, leftMargin=20*mm, topMargin=20*mm, bottomMargin=20*mm)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle('title', fontSize=18, fontName='Helvetica-Bold',
                                      textColor=colors.HexColor('#FF6B00'), alignment=TA_CENTER)
        story.append(Paragraph(f"RAPPORT FINANCIER {year}", title_style))
        story.append(Spacer(1, 8*mm))

        # Monthly breakdown
        data = [["Mois", "Recettes", "Dépenses", "Bénéfice"]]
        total_rev = 0
        total_exp = 0
        for m in range(1, 13):
            rev = self.session.query(func.sum(Payment.amount)).filter(
                Payment.month == m, Payment.year == year).scalar() or 0
            exp = self.session.query(func.sum(Expense.amount)).filter(
                func.strftime('%m', Expense.expense_date) == f"{m:02d}",
                func.strftime('%Y', Expense.expense_date) == str(year)).scalar() or 0
            profit = rev - exp
            data.append([MONTHS[m-1], f"{rev:,.2f} MAD", f"{exp:,.2f} MAD", f"{profit:,.2f} MAD"])
            total_rev += rev
            total_exp += exp

        data.append(["TOTAL", f"{total_rev:,.2f} MAD", f"{total_exp:,.2f} MAD", f"{total_rev-total_exp:,.2f} MAD"])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#FF6B00')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
            ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#1a1a2e')),
            ('TEXTCOLOR', (0,-1), (-1,-1), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#FFF5EE')]),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e0e0e0')),
            ('TOPPADDING', (0,0), (-1,-1), 7),
            ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ]))
        story.append(table)
        doc.build(story)
        self._open_file(filepath)
        QMessageBox.information(self, "Succès", f"✅ Rapport financier généré!")

    def _generate_unpaid_pdf(self):
        from models.database import Student, Payment
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.lib.enums import TA_CENTER

        month = self.month_combo.currentIndex() + 1
        year = self.year_spin.value()
        paid_ids = [p[0] for p in self.session.query(Payment.student_id).filter(
            Payment.month == month, Payment.year == year, Payment.payment_type == 'monthly').all()]
        unpaid = self.session.query(Student).filter(
            Student.is_active == True, ~Student.id.in_(paid_ids)).all()

        reports_dir = os.path.join(BASE_DIR, 'reports')
        os.makedirs(reports_dir, exist_ok=True)
        filepath = os.path.join(reports_dir, f'unpaid_{year}_{month:02d}.pdf')

        doc = SimpleDocTemplate(filepath, pagesize=A4,
                               rightMargin=20*mm, leftMargin=20*mm, topMargin=20*mm, bottomMargin=20*mm)
        styles = getSampleStyleSheet()
        story = []
        title_style = ParagraphStyle('title', fontSize=16, fontName='Helvetica-Bold',
                                      textColor=colors.HexColor('#e74c3c'), alignment=TA_CENTER)
        story.append(Paragraph(f"ÉLÈVES IMPAYÉS - {MONTHS[month-1].upper()} {year}", title_style))
        story.append(Spacer(1, 8*mm))

        data = [["Nom", "Prénom", "Classe", "Frais mensuel", "Parent", "Téléphone"]]
        total = 0
        for s in unpaid:
            data.append([s.last_name, s.first_name,
                        s.class_.name if s.class_ else "—",
                        f"{s.monthly_fee:,.2f} MAD",
                        s.parent_name or "—", s.parent_phone or "—"])
            total += s.monthly_fee or 0

        if len(data) == 1:
            story.append(Paragraph("✅ Tous les élèves sont à jour!", styles['Normal']))
        else:
            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e74c3c')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#FFF5EE')]),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e0e0e0')),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ]))
            story.append(table)
            story.append(Spacer(1, 5*mm))
            story.append(Paragraph(f"Total impayé: {total:,.2f} MAD | {len(unpaid)} élève(s)", styles['Normal']))

        doc.build(story)
        self._open_file(filepath)
        QMessageBox.information(self, "Succès", f"✅ Rapport des impayés généré!")

    def _generate_employees_pdf(self):
        from models.database import Employee
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.lib.enums import TA_CENTER

        reports_dir = os.path.join(BASE_DIR, 'reports')
        os.makedirs(reports_dir, exist_ok=True)
        filepath = os.path.join(reports_dir, f'employees_{datetime.now().strftime("%Y%m%d")}.pdf')

        doc = SimpleDocTemplate(filepath, pagesize=A4,
                               rightMargin=20*mm, leftMargin=20*mm, topMargin=20*mm, bottomMargin=20*mm)
        story = []
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('title', fontSize=16, fontName='Helvetica-Bold',
                                      textColor=colors.HexColor('#9b59b6'), alignment=TA_CENTER)
        story.append(Paragraph("RAPPORT DES EMPLOYÉS", title_style))
        story.append(Spacer(1, 8*mm))

        employees = self.session.query(Employee).filter_by(is_active=True).all()
        data = [["Nom", "Prénom", "Type", "Matière", "Téléphone", "Salaire"]]
        total_salary = 0
        type_map = {'teacher': 'Enseignant', 'staff': 'Personnel', 'driver': 'Chauffeur'}
        for emp in employees:
            data.append([emp.last_name, emp.first_name,
                        type_map.get(emp.employee_type, emp.employee_type),
                        emp.subject or "—", emp.phone or "—",
                        f"{emp.base_salary:,.0f} MAD"])
            total_salary += emp.base_salary or 0

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#9b59b6')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#FFF5EE')]),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e0e0e0')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(table)
        story.append(Spacer(1, 5*mm))
        story.append(Paragraph(f"Total salaires: {total_salary:,.0f} MAD/mois | {len(employees)} employé(s)", styles['Normal']))
        doc.build(story)
        self._open_file(filepath)
        QMessageBox.information(self, "Succès", f"✅ Rapport des employés généré!")

    def export_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from models.database import Student, Payment
            from PySide6.QtWidgets import QFileDialog

            reports_dir = os.path.join(BASE_DIR, 'reports')
            os.makedirs(reports_dir, exist_ok=True)
            default_path = os.path.join(reports_dir, f'export_{datetime.now().strftime("%Y%m%d")}.xlsx')
            filepath, _ = QFileDialog.getSaveFileName(self, "Enregistrer Excel", default_path, "Excel (*.xlsx)")
            if not filepath:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Élèves"

            orange = "FFFF6B00"
            headers = ["Code", "Nom", "Prénom", "Classe", "Parent", "Téléphone", "Frais Mensuel", "Transport", "Assurance"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFFFF")
                cell.fill = PatternFill("solid", fgColor=orange)
                cell.alignment = Alignment(horizontal='center')

            students = self.session.query(Student).filter_by(is_active=True).all()
            for row_idx, s in enumerate(students, 2):
                row_data = [
                    s.student_code, s.last_name, s.first_name,
                    s.class_.name if s.class_ else "",
                    s.parent_name, s.parent_phone, s.monthly_fee,
                    "Oui" if s.has_transport else "Non",
                    "Oui" if s.insurance_paid else "Non"
                ]
                for col_idx, val in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=val)

            for column in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in column)
                ws.column_dimensions[column[0].column_letter].width = min(max_len + 3, 25)

            # Payments sheet
            ws2 = wb.create_sheet("Paiements")
            headers2 = ["Reçu", "Élève", "Type", "Mois", "Année", "Montant", "Date"]
            for col, h in enumerate(headers2, 1):
                cell = ws2.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFFFF")
                cell.fill = PatternFill("solid", fgColor=orange)
                cell.alignment = Alignment(horizontal='center')

            type_map = {'monthly': 'Mensualité', 'insurance': 'Assurance', 'transport': 'Transport'}
            payments = self.session.query(Payment).order_by(Payment.payment_date.desc()).limit(500).all()
            for row_idx, p in enumerate(payments, 2):
                student = p.student
                row_data = [
                    p.receipt_number,
                    f"{student.first_name} {student.last_name}" if student else "",
                    type_map.get(p.payment_type, p.payment_type),
                    MONTHS[p.month-1] if p.month and 1 <= p.month <= 12 else "",
                    str(p.year or ""), p.amount,
                    p.payment_date.strftime("%d/%m/%Y") if p.payment_date else ""
                ]
                for col_idx, val in enumerate(row_data, 1):
                    ws2.cell(row=row_idx, column=col_idx, value=val)

            wb.save(filepath)
            self._open_file(filepath)
            QMessageBox.information(self, "Succès", f"✅ Export Excel généré!")
        except Exception as e:
            QMessageBox.warning(self, "Erreur", f"Erreur lors de l'export:\n{e}")

    def _open_file(self, filepath):
        import subprocess, platform
        try:
            if platform.system() == 'Windows':
                os.startfile(filepath)
            elif platform.system() == 'Darwin':
                subprocess.run(['open', filepath])
            else:
                subprocess.run(['xdg-open', filepath])
        except:
            pass
